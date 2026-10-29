"""CSV / Excel / pinned-host SFTP connectors with tenant-scoped paths and bounded reads."""

import csv
import hashlib
import io
import posixpath
from datetime import UTC, datetime
from pathlib import PurePosixPath

from openpyxl import load_workbook
from paramiko import RejectPolicy, SSHClient

from ..service import DomainError
from .connector_sdk import Connector, json_value
from .contracts import ReadBatch, SourceColumn, TableSchema

MAX_BYTES = 20 * 1024 * 1024
MAX_FILE_ROWS = 100000


def load_tables(content, filename):
    if len(content) > MAX_BYTES:
        raise DomainError("来源文件超过 20 MiB 上限")
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix == ".csv":
        try:
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            matrix = []
            for index, row in enumerate(reader):
                if index > MAX_FILE_ROWS or len(row) > 256:
                    raise DomainError("文件超过 100,000 行或 256 列上限")
                matrix.append(row)
            tables = {PurePosixPath(filename).stem: matrix}
        except (UnicodeError, csv.Error) as exc:
            raise DomainError("CSV 编码或内容不合法") from exc
    elif suffix == ".xlsx":
        # Bound uncompressed size as well as zip file size (zip-bomb protection).
        from zipfile import ZipFile

        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > MAX_BYTES * 10:
                    raise DomainError("Excel 解压后超过大小上限")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > 50:
                    raise DomainError("Excel 工作表超过 50 个")
                tables = {}
                for sheet in workbook.worksheets:
                    if sheet.max_column and sheet.max_column > 256:
                        raise DomainError("Excel 超过 256 列上限")
                    values = []
                    for index, row in enumerate(sheet.iter_rows()):
                        if index > MAX_FILE_ROWS:
                            raise DomainError("文件超过 100,000 行上限")
                        if any(cell.data_type == "f" for cell in row):
                            raise DomainError("Excel 包含公式，请先导出为确定的值；不会使用未验证缓存")
                        values.append([json_value(cell.value) for cell in row])
                    tables[sheet.title] = values
            finally:
                workbook.close()
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError("Excel 文件无法解析") from exc
    else:
        raise DomainError("只支持 UTF-8 CSV 或 XLSX 文件")
    result = {}
    for table, matrix in tables.items():
        if not matrix:
            raise DomainError("工作表缺少表头")
        header = matrix[0]
        if not all(isinstance(c, str) and c.strip() for c in header) or len(header) != len(set(header)):
            raise DomainError("表头必须为唯一的非空文本")
        rows = []
        for row in matrix[1:]:
            if len(row) != len(header):
                raise DomainError("数据行的列数与表头不一致")
            if all(value is None for value in row):
                continue
            rows.append(dict(zip(header, row, strict=True)))
            if len(rows) > MAX_FILE_ROWS:
                raise DomainError("文件超过 100,000 行上限")
        result[table] = (header, rows)
    return result


class FileConnector(Connector):
    kinds = ("csv", "excel")

    def content(self):
        path = self.context.path()
        before = path.stat()
        if before.st_size > MAX_BYTES:
            raise DomainError("文件超过 20 MiB 上限")
        content = path.read_bytes()
        after = path.stat()
        if before.st_mtime_ns != after.st_mtime_ns or before.st_size != after.st_size:
            raise DomainError("来源文件正在变化，请在导出完成后重试", 409)
        return content, path.name, datetime.fromtimestamp(after.st_mtime, UTC)

    def tables(self):
        content, name, modified = self.content()
        return load_tables(content, name), modified, hashlib.sha256(content).hexdigest()

    def discover(self):
        tables, _, _ = self.tables()
        result = []
        for name, (header, rows) in tables.items():
            columns = []
            for column in header:
                values = [row[column] for row in rows if row[column] is not None]
                candidate = (
                    bool(values) and len(values) == len(rows) and len({str(v) for v in values}) == len(values)
                )
                types = {type(v).__name__ for v in values[:100]}
                columns.append(
                    SourceColumn(
                        name=column,
                        type="/".join(sorted(types)) or "unknown",
                        nullable=len(values) != len(rows),
                        primary_key_candidate=candidate,
                    )
                )
            result.append(TableSchema(name=name, columns=columns, sample=rows[:3]))
        return result

    def read(self, table, primary_key, batch_size, cursor=None, checkpoint=None):
        if cursor:
            raise DomainError("文件来源使用带内容指纹的全量快照，不接受数据库增量游标")
        tables, modified, fingerprint = self.tables()
        if table not in tables:
            raise DomainError("工作表不存在")
        header, rows = tables[table]
        if primary_key not in header:
            raise DomainError("主键不在文件表头中")
        keys = [row[primary_key] for row in rows]
        if any(key is None or str(key) == "" for key in keys) or len(set(map(str, keys))) != len(keys):
            raise DomainError("文件主键为空或重复")
        if checkpoint and checkpoint.get("fingerprint") != fingerprint:
            raise DomainError("来源文件已变化，本批次必须从头重试", 409)
        start = (checkpoint or {}).get("offset", 0)
        end = min(start + batch_size, len(rows))
        return ReadBatch(
            rows=rows[start:end],
            checkpoint={"offset": end, "fingerprint": fingerprint},
            exhausted=end == len(rows),
            source_updated_at=modified,
            freshness_basis="source_file_mtime",
        )


class SFTPConnector(FileConnector):
    kinds = ("sftp",)

    def content(self):
        config = self.context.profile()
        remote = self.context.source.get("path")
        root = posixpath.normpath(config["root"])
        if not root.startswith("/") or not remote:
            raise DomainError("SFTP 需要绝对授权根目录和相对来源路径")
        path = posixpath.normpath(posixpath.join(root, remote))
        if not path.startswith(root.rstrip("/") + "/"):
            raise DomainError("SFTP 路径超出授权根目录")
        client = SSHClient()
        try:
            client.load_host_keys(config["known_hosts"])
            client.set_missing_host_key_policy(RejectPolicy())
            client.connect(
                hostname=config["host"],
                port=config.get("port", 22),
                username=config["username"],
                password=config.get("password"),
                key_filename=config.get("key_filename"),
                timeout=5,
                auth_timeout=5,
                banner_timeout=5,
                allow_agent=False,
                look_for_keys=False,
            )
            with client.open_sftp() as sftp:
                sftp.get_channel().settimeout(15)
                real = sftp.normalize(path)
                if not real.startswith(root.rstrip("/") + "/"):
                    raise DomainError("SFTP 符号链接超出授权根目录")
                before = sftp.stat(real)
                if before.st_size > MAX_BYTES:
                    raise DomainError("SFTP 文件超过 20 MiB")
                with sftp.open(real, "rb") as stream:
                    content = stream.read(MAX_BYTES + 1)
                after = sftp.stat(real)
                if before.st_mtime != after.st_mtime or before.st_size != after.st_size:
                    raise DomainError("远程文件正在变化，请重试", 409)
                return content, PurePosixPath(real).name, datetime.fromtimestamp(after.st_mtime, UTC)
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError("SFTP 连接失败，请检查固定主机指纹、授权路径和凭据", 502) from exc
        finally:
            client.close()
